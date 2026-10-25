from contextlib import closing
from pathlib import Path
import json
import sqlite3
import tempfile
import unittest
from unittest.mock import patch

CAPTURE='2026-10-22T08:30:00+09:00'


class PublicReportingTests(unittest.TestCase):
    def test_stock_percent_uses_same_six_decimal_contract_as_market_table(self):
        import fdr_provider,report_model,FastmarketsCollector_v3_4 as core
        with tempfile.TemporaryDirectory() as t:
            fdr_provider.collect_and_save(core.path_bundle(Path(t))['db'],lambda:False,lambda _:None,
                captured_at=CAPTURE,fetcher=lambda *a:[dict(date='2026-10-20',close='143'),dict(date='2026-10-21',close='144')])
            c=report_model.build_context(t,captured_at=CAPTURE)
            self.assertEqual(c.tables['보고용_주가'][1][7],round(100/143,6))

    def fixture(self,folder,missing=False):
        import FastmarketsCollector_v3_4 as core
        import fdr_provider
        db=core.path_bundle(Path(folder))['db']
        def fetch(spec,*args):
            days=['2026-10-19','2026-10-20'] if missing and spec.code=='KOSPI' else ['2026-10-20','2026-10-21']
            return [dict(date=d,close=v) for d,v in zip(days,['100','101'])]
        fdr_provider.collect_and_save(db,lambda:False,lambda _:None,captured_at=CAPTURE,fetcher=fetch)
        return db

    def test_db_to_report_expected_dates_source_and_no_fabricated_rank(self):
        import report_model
        with tempfile.TemporaryDirectory() as t:
            self.fixture(t)
            c=report_model.build_context(t,end='2026-10-22',captured_at=CAPTURE)
            row=dict(zip(c.tables['보고용_시장지표'][0],c.tables['보고용_시장지표'][1]))
            self.assertEqual(row['수치'],101)
            self.assertEqual(row['시장일'],'2026-10-21')
            self.assertEqual(row['비교일'],'2026-10-20')
            self.assertIn('finance.naver.com',row['출처'])
            self.assertEqual(c.tables['보고용_주가'][1][0],'005490')
            self.assertEqual(len(c.tables['보고용_시총']),1)
            sheets=report_model.make_workbook_sheets(c)
            self.assertIn('FDR_종가원본',sheets)
            self.assertTrue(any(str(r[0]).startswith('FDR|KOSPI|') for r in sheets['가격열설정'][1:]))
            self.assertFalse(any(str(r[0]).startswith('NAVER|KOSPI|') for r in sheets['가격열설정'][1:]))

    def test_missing_latest_not_reported_as_older_close_pair(self):
        import report_model
        with tempfile.TemporaryDirectory() as t:
            self.fixture(t,missing=True)
            c=report_model.build_context(t,end='2026-10-22',captured_at=CAPTURE)
            row=dict(zip(c.tables['보고용_시장지표'][0],c.tables['보고용_시장지표'][1]))
            self.assertEqual(row['상태'],'missing_history')
            self.assertIsNone(row['수치'])
            self.assertEqual(row['시장일'],'2026-10-21')
            self.assertIsNone(row['전일대비'])

    def test_cross_format_native_table_and_snapshot_consistency(self):
        import report_model,report_exports
        from pptx import Presentation
        from openpyxl import load_workbook
        with tempfile.TemporaryDirectory() as t:
            self.fixture(t)
            c=report_model.build_context(t,end='2026-10-22',synthetic=True,captured_at=CAPTURE)
            result=report_exports.export_reports(c,Path(t)/'out',include_excel=True)
            html=Path(result['html']).read_text()
            self.assertIn('finance.yahoo.com',html)
            self.assertIn(c.snapshot_id,html)
            workbook=load_workbook(result['excel'],read_only=True)
            try:
                self.assertEqual(workbook['보고용_주가']['A2'].value,'005490')
                for name,rows in c.tables.items():
                    if name=='보고용_종합':continue
                    # XLSX serializes empty text as blank; preserve every other value.
                    expected=[[None if v=='' else v for v in row] for row in rows]
                    self.assertEqual([list(r) for r in workbook[name].values],expected,name)
            finally:workbook.close()
            prs=Presentation(result['pptx'])
            cells=[cell.text for slide in prs.slides for shape in slide.shapes if shape.has_table for row in shape.table.rows for cell in row.cells]
            self.assertIn('005490',cells)
            for name,rows in c.tables.items():
                if name=='보고용_종합':continue
                for row in rows:
                    for value in row:self.assertIn(report_exports.display_value(value),cells)
            self.assertEqual(json.loads(Path(result['context']).read_text())['snapshot_id'],c.snapshot_id)

    def test_full_series_filter_retains_raw_provenance_and_only_its_dates(self):
        import report_model
        with tempfile.TemporaryDirectory() as t:
            self.fixture(t)
            c=report_model.build_context(t,series=['FDR|KOSPI|points|Actual|MID'],captured_at=CAPTURE)
            self.assertGreater(len(c._sheets['FDR_종가원본']),1)
            rows=c._sheets['FDR_종가원본'];i=rows[0].index('instrument_id')
            self.assertEqual({row[i] for row in rows[1:]},{'KOSPI'})

    def test_full_series_filter_initial_failed_fetch_still_shows_missing(self):
        import fdr_provider,report_model,FastmarketsCollector_v3_4 as core
        with tempfile.TemporaryDirectory() as t:
            db=core.path_bundle(Path(t))['db']
            fdr_provider.collect_and_save(db,lambda:False,lambda _:None,captured_at=CAPTURE,fetcher=lambda *a:[])
            c=report_model.build_context(t,series=['FDR|KOSPI|points|Actual|MID'],captured_at=CAPTURE)
            self.assertEqual(c.tables['보고용_시장지표'][1][0],'KOSPI')
            self.assertEqual(c.tables['보고용_시장지표'][1][8],'missing_history')

    def test_stock_only_has_actual_capture_as_of(self):
        import report_model
        with tempfile.TemporaryDirectory() as t:
            self.fixture(t)
            c=report_model.build_context(t,series=['005490'],captured_at=CAPTURE)
            self.assertEqual(c.as_of,CAPTURE)

    def test_historical_stock_registry_can_exceed_per_request_limit(self):
        import report_model,fdr_provider,FastmarketsCollector_v3_4 as core
        with tempfile.TemporaryDirectory() as t:
            db=core.path_bundle(Path(t))['db']
            fetch=lambda *a:[{'date':'2026-10-20','close':'100'},{'date':'2026-10-21','close':'101'}]
            for codes in ([f'{i:06}' for i in range(25)],['005490']):
                fdr_provider.collect_and_save(db,lambda:False,lambda _:None,stock_codes=codes,captured_at=CAPTURE,fetcher=fetch)
            c=report_model.build_context(t,series=['005490'],captured_at=CAPTURE)
            self.assertEqual(len(c.tables['보고용_주가']),2)
            self.assertEqual(c.tables['보고용_주가'][1][0],'005490')

    def test_refresh_conflict_withdraws_cached_close_and_valid_refresh_restores(self):
        import report_model,fdr_provider
        with tempfile.TemporaryDirectory() as t:
            db=self.fixture(t)
            before=report_model.make_workbook_sheets(report_model.build_context(t,captured_at=CAPTURE))
            def conflicted(spec,*args):
                return [{'date':'2026-10-20','close':'100'},{'date':'2026-10-21','close':'101'},{'date':'2026-10-21','close':'999'}]
            fdr_provider.collect_and_save(db,lambda:False,lambda _:None,captured_at=CAPTURE,fetcher=conflicted,force=True)
            c=report_model.build_context(t,captured_at=CAPTURE)
            self.assertIsNone(c.tables['보고용_시장지표'][1][2])
            self.assertEqual(c.tables['보고용_시장지표'][1][8],'missing_history')
            self.assertIn('FDR_보류이력',c._sheets)
            from excel_secure import quarantined_cells
            from ecm_excel_sync import build_plan
            after=report_model.make_workbook_sheets(c)
            plan=build_plan(after['일별시세'],after['가격열설정'],before['일별시세'],before['가격열설정'],
                clear_cells=quarantined_cells(after,'일별시세'))
            self.assertEqual(plan['stats']['cleared'],7)
            with closing(sqlite3.connect(db)) as cx:
                self.assertEqual(cx.execute("SELECT value FROM fdr_observations WHERE instrument_id='KOSPI' AND market_date='2026-10-21'").fetchone()[0],'101')
            fdr_provider.collect_and_save(db,lambda:False,lambda _:None,captured_at=CAPTURE,
                fetcher=lambda *a:[{'date':'2026-10-20','close':'100'},{'date':'2026-10-21','close':'101'}],force=True)
            restored=report_model.build_context(t,captured_at=CAPTURE)
            self.assertEqual(restored.tables['보고용_시장지표'][1][2],101)
            self.assertEqual(restored.tables['보고용_시장지표'][1][8],'close_pair')


if __name__=='__main__':unittest.main()
