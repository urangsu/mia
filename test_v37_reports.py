"""Synthetic-only cross-format snapshot contracts."""
import json
import sqlite3
import tempfile
import unittest
from contextlib import closing
from pathlib import Path
from unittest.mock import patch

class ReportTests(unittest.TestCase):
    def test_stale_without_close_evidence_not_daily_or_caps_all_expected_missing_visible(self):
        import report_model as m
        from naver_markets import INDICATORS
        with tempfile.TemporaryDirectory() as t:
            db=self.fixture(t)
            with closing(sqlite3.connect(db)) as cx,cx:
                cx.execute("UPDATE market_indicators SET status='stale',raw_json='{}'")
                cx.execute("UPDATE equity_rank_snapshots SET status='stale'")
            c=m.build_context(t)
            self.assertEqual(len(c.tables['보고용_시장지표'])-1,len(INDICATORS))
            self.assertTrue(all(r[2] is None and r[8]=='missing_history' for r in c.tables['보고용_시장지표'][1:]))
            self.assertEqual(len(c.tables['보고용_시총']),1)
            self.assertNotIn('2500.25',str(m.make_workbook_sheets(c)['일별시세']))
            self.assertEqual(len(m.make_workbook_sheets(c)['NAVER_시총원본']),2)

    def test_naver_report_uses_two_closed_days_not_vendor_intraday_delta(self):
        import report_model as m
        with tempfile.TemporaryDirectory() as t:
            db=self.fixture(t)
            with closing(sqlite3.connect(db)) as cx,cx:
                for day,value,status in [('2026-10-15','2400.25','final'),('2026-10-19','9999','provisional')]:
                    cx.execute('INSERT INTO market_indicators VALUES(?,?,?,?,?,?,?,?,?,?,?)',('KOSPI',day,'2026-10-19T08:30:00+09:00',value,'777','88','point','',status,'https://finance.naver.com/','{}'))
                cx.execute('INSERT INTO equity_rank_snapshots VALUES(?,?,?,?,?,?,?,?,?,?,?,?,?,?)',('TODAY',1,'장중','2026-10-19','2026-10-19T08:30:00+09:00',1000,1,1,999,'원','KRW','provisional','',json.dumps({'market_cap_unit':'억원'})))
            c=m.build_context(t,end='2026-10-19')
            row=dict(zip(c.tables['보고용_시장지표'][0],c.tables['보고용_시장지표'][1]))
            self.assertEqual(row['시장일'],'2026-10-16')
            self.assertEqual(row['수치'],2500.25)
            self.assertEqual(row['비교일'],'2026-10-15')
            self.assertEqual(row['비교종가'],2400.25)
            self.assertEqual(row['전일대비'],100)
            self.assertEqual(row['상태'],'close_pair')
            self.assertNotIn('TODAY',str(c.tables['보고용_시총']))
            self.assertNotIn('9999',str(m.make_workbook_sheets(c)['일별시세']))
            missing=m.build_context(t,start='2026-10-16',end='2026-10-19')
            row=dict(zip(missing.tables['보고용_시장지표'][0],missing.tables['보고용_시장지표'][1]))
            self.assertIsNone(row['비교종가']);self.assertIsNone(row['전일대비'])
            self.assertEqual(row['상태'],'missing_history')
            historical=m.build_context(t,end='2026-10-16')
            self.assertEqual(historical.tables['보고용_시장지표'][1][1],'2026-10-15')

    def test_unchanged_db_produces_stable_snapshot_and_workbook_digest(self):
        import report_model as m
        from export_cache import digest_sheets
        with tempfile.TemporaryDirectory() as t:
            self.fixture(t)
            a=m.build_context(t)
            b=m.build_context(t)
            self.assertEqual(a.as_of,b.as_of)
            self.assertEqual(a.snapshot_id,b.snapshot_id)
            self.assertEqual(digest_sheets(m.make_workbook_sheets(a)),digest_sheets(m.make_workbook_sheets(b)))
            self.assertEqual(a.as_of,'2026-10-19T09:00:00+09:00')

    def test_fx_duplicate_sources_quarantine_conflict_preserve_raw(self):
        import report_model as m
        from smbs_rates import FXQuote,prepare_excel_series
        quotes=[FXQuote('2026-10-16','USD','1400','1',1,'SMBS'),FXQuote('2026-10-16','USD','999','1',1,'legacy')]
        issues=[];values=prepare_excel_series([],quotes,issues=issues)
        self.assertNotIn('FX_USD_KRW_per_1',values.get('2026-10-16',{}))
        self.assertTrue(any('충돌' in w for w in issues))
        with tempfile.TemporaryDirectory() as t:
            db=self.fixture(t)
            with closing(sqlite3.connect(db)) as cx,cx:
                cx.execute('CREATE TABLE fx_rates(source,currency,quote_date,quote_unit,rate,difference,collected_at)')
                cx.executemany('INSERT INTO fx_rates VALUES(?,?,?,?,?,?,?)',[(q.source,q.currency,q.quote_date,q.unit,q.rate,q.difference,'2026-10-19') for q in quotes])
            c=m.build_context(t)
            self.assertEqual(len(c.tables['보고용_환율']),2)
            row=c.tables['보고용_환율'][1]
            self.assertIsNone(row[2]);self.assertIsNone(row[3])
            self.assertIn('legacy',row[5]);self.assertIn('SMBS',row[5])
            self.assertEqual(len(m.make_workbook_sheets(c)['환율원본']),3)

    def test_fx_duplicate_sources_same_values_allowed_diff_conflict_quarantined(self):
        from smbs_rates import FXQuote,prepare_excel_series
        a=FXQuote('2026-10-16','USD','1400','1',1,'SMBS')
        b=FXQuote('2026-10-16','USD','1400.0','1.0',1,'legacy')
        issues=[];r=prepare_excel_series([],[a,b],issues=issues)
        self.assertEqual(r['2026-10-16']['FX_USD_KRW_per_1'],1400)
        self.assertFalse(issues)
        c=FXQuote('2026-10-16','USD','1400','2',1,'other')
        issues=[];r=prepare_excel_series([],[a,c],issues=issues)
        self.assertNotIn('FX_USD_KRW_per_1',r.get('2026-10-16',{}))
        self.assertNotIn('FX_USD_DIFF_KRW_per_1',r.get('2026-10-16',{}))

    def test_numeric_display_does_not_round_small_values_to_zero(self):
        from report_exports import display_value
        self.assertEqual(display_value(0.000000000001),'0.000000000001')
        self.assertEqual(display_value(1000),'1,000')
        from report_model import _number
        self.assertIsNone(_number('1e9999'))

    def fixture(self, folder):
        import FastmarketsCollector_v3_4 as core
        from naver_markets import initialize_schema
        db=core.path_bundle(Path(folder))['db']
        db.parent.mkdir(parents=True,exist_ok=True)
        with closing(sqlite3.connect(db)) as cx,cx:
            initialize_schema(cx)
            cx.execute("INSERT INTO market_indicators VALUES(?,?,?,?,?,?,?,?,?,?,?)",('KOSPI','2026-10-16','2026-10-19T09:00:00+09:00','2500.25','25','1.01','point','','final','https://finance.naver.com/','{}'))
            cx.execute("INSERT INTO market_indicators VALUES(?,?,?,?,?,?,?,?,?,?,?)",('BAD','2026-10-16','2026-10-19','999',None,None,'point','','unverified','','{}'))
            cx.execute("INSERT INTO equity_rank_snapshots VALUES(?,?,?,?,?,?,?,?,?,?,?,?,?,?)",('005930',1,'=BAD','2026-10-16','2026-10-19',70000,100,0.14,400000,'원','KRW','final','',json.dumps({'market_cap_unit':'억원'})))
        return db

    def test_context_immutable_filtered_and_no_equity_daily(self):
        import report_model as m
        with tempfile.TemporaryDirectory() as t:
            self.fixture(t)
            c=m.build_context(t,end='2026-10-17',synthetic=True)
            self.assertIn('보고용_금속시세',c.tables)
            self.assertNotIn('BAD',str(c.tables['보고용_시장지표']))
            self.assertIn('2026-10-16',str(c.source_dates))
            with self.assertRaises(TypeError):c.tables['other']=()
            sheets=m.make_workbook_sheets(c)
            self.assertTrue(any(str(r[0]).startswith('NAVER|KOSPI|') for r in sheets['가격열설정'][1:]))
            self.assertNotIn('005930',str(sheets['일별시세']))
            self.assertEqual(sheets['보고용_시총'][1][1],"'=BAD")
            self.assertIn('TEST / 합성 데이터',c.label)
            early=m.build_context(t,end='2026-10-15')
            self.assertTrue(all(r[2] is None for r in early.tables['보고용_시장지표'][1:]))

    def test_exports_exact_same_cells_and_frozen_after_db_change(self):
        import report_model as m
        from report_exports import export_reports,display_value
        from pptx import Presentation
        from openpyxl import load_workbook
        with tempfile.TemporaryDirectory() as t:
            db=self.fixture(t)
            c=m.build_context(t,synthetic=True)
            with closing(sqlite3.connect(db)) as cx,cx:cx.execute("UPDATE market_indicators SET value='9999'")
            result=export_reports(c,Path(t)/'reports',include_excel=True)
            html=Path(result['html']).read_text()
            self.assertIn('2,500.25',html)
            self.assertNotIn('9,999',html)
            self.assertIn(c.snapshot_id,html)
            self.assertIn('<script src="market_report.js" defer></script>',html)
            self.assertTrue((Path(result['html']).parent / 'market_report.js').is_file())
            self.assertIn('자료 기준 · 출처 · 검증 상태 펼치기',html)
            self.assertNotIn('<script>alert(',html)
            prs=Presentation(result['pptx'])
            self.assertGreaterEqual(len(prs.slides),3)
            native=[cell.text for slide in prs.slides for shape in slide.shapes if shape.has_table for row in shape.table.rows for cell in row.cells]
            for name,rows in c.tables.items():
                if name=='보고용_종합':continue
                for row in rows:
                    for cell in row:self.assertIn(display_value(cell),native)
            wb=load_workbook(result['excel'],data_only=False)
            self.assertEqual(wb['보고서_메타']['B2'].value,c.snapshot_id)
            self.assertEqual(wb['보고용_시장지표']['C2'].value,2500.25)
            self.assertEqual(wb['보고용_시총']['B2'].data_type,'s')
            evidence=json.loads(Path(result['context']).read_text())
            self.assertEqual(evidence['snapshot_id'],c.snapshot_id)
            from report_model import canonical_json
            from hashlib import sha256
            evidence.pop('snapshot_id')
            self.assertEqual(sha256(canonical_json(evidence).encode()).hexdigest(),c.snapshot_id)
            wb.close()

    def test_conflicts_missing_mid_and_other_workspace_quarantined(self):
        import report_model as m
        with tempfile.TemporaryDirectory() as t:
            db=self.fixture(t)
            with closing(sqlite3.connect(db)) as cx,cx:
                cx.execute('CREATE TABLE prices(workspace_id,symbol,assessment_date,low,mid,high,change_pct,unit,description,widget,price_mode,source,view_key,period_label,collected_at)')
                for workspace,code,mid,view in [('active','MB-X','15','one'),('active','MB-X','99','two'),('active','MB-MISSING',None,'three'),('other','MB-X','55','other')]:
                    cx.execute('INSERT INTO prices VALUES(?,?,?,?,?,?,?,?,?,?,?,?,?,?,?)',(workspace,code,'2026-10-16','10',mid,'20',None,'USD/t','Synthetic','test','Actual','FMB',view,'','2026-10-19'))
            c=m.build_context(t,workspace_id='active')
            self.assertTrue(all(r[2] is None for r in c.tables['보고용_금속시세'][1:]))
            self.assertTrue(any('충돌' in w for w in c.warnings))
            self.assertEqual(len(c._sheets['워크스페이스원본']),5)
            selected=m.build_context(t,workspace_id='active',series=['MB-MISSING'])
            self.assertEqual(len(selected.tables['보고용_금속시세']),2)
            self.assertEqual(len(selected.tables['보고용_시장지표']),1)

    def test_pptx_paginates_all_rows_and_html_escapes(self):
        import report_model as m
        from report_exports import export_reports
        from pptx import Presentation
        with tempfile.TemporaryDirectory() as t:
            db=self.fixture(t)
            with closing(sqlite3.connect(db)) as cx,cx:
                cx.execute("UPDATE equity_rank_snapshots SET name='<script>alert(1)</script>'")
                for i in range(26):
                    cx.execute('INSERT INTO market_indicators VALUES(?,?,?,?,?,?,?,?,?,?,?)',('SYNTH'+str(i),'2026-10-16','2026-10-19',str(i+0.125),None,None,'point','','stale','',json.dumps({'close_verified':True})))
            c=m.build_context(t,synthetic=True)
            result=export_reports(c,Path(t)/'out')
            html=Path(result['html']).read_text()
            self.assertIn('&lt;script&gt;',html)
            self.assertNotIn('<script>',html)
            p=Presentation(result['pptx']);seen=[]
            for slide in p.slides:
                for shape in slide.shapes:
                    if shape.has_table:
                        self.assertLessEqual(len(shape.table.rows),13)
                        seen.extend(cell.text for row in shape.table.rows for cell in row.cells)
            for i in range(26):self.assertIn('SYNTH'+str(i),seen)

    def test_capitalization_never_reuses_price_unit(self):
        import report_model as m
        with tempfile.TemporaryDirectory() as t:
            db=self.fixture(t)
            with closing(sqlite3.connect(db)) as cx,cx:cx.execute("UPDATE equity_rank_snapshots SET raw_json='{}'")
            c=m.build_context(t)
            row=c.tables['보고용_시총'][1]
            self.assertIsNone(row[4])
            self.assertEqual(row[5],'미확인')
            self.assertTrue(any('시총 단위' in w for w in c.warnings))
            from report_exports import export_reports
            from pptx import Presentation
            result=export_reports(c,Path(t)/'out')
            native=' '.join(shape.text for slide in Presentation(result['pptx']).slides for shape in slide.shapes if shape.has_text_frame)
            self.assertIn('시총 단위',native)

    def test_posco_extra_keeps_own_date_without_mixing_ranked_dates(self):
        import report_model as m
        with tempfile.TemporaryDirectory() as t:
            db=self.fixture(t)
            with closing(sqlite3.connect(db)) as cx,cx:
                for code,rank,day,cap in [('005490',None,'2026-10-15','10000'),('005490',None,'2026-10-14','9000'),('OLD',2,'2026-10-14','8000')]:
                    cx.execute('INSERT INTO equity_rank_snapshots VALUES(?,?,?,?,?,?,?,?,?,?,?,?,?,?)',(code,rank,'Synthetic',day,'2026-10-19',1000,0,0,cap,'원','KRW','stale','',json.dumps({'market_cap_unit':'억원','close_verified':True})))
            c=m.build_context(t)
            rows=c.tables['보고용_시총'][1:]
            self.assertEqual(len(rows),2)
            posco=next(r for r in rows if r[2]=='005490')
            self.assertIsNone(posco[0])
            self.assertEqual(posco[3],'2026-10-15')
            self.assertEqual(posco[4],10000)
            self.assertNotIn('OLD',str(rows))

if __name__=='__main__':unittest.main()
