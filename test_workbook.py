import tempfile
import unittest
from datetime import date
from pathlib import Path
from unittest.mock import patch
from test_v33 import m, price


class WorkbookTests(unittest.TestCase):
    def test_custom_workbook_path_preserves_notes(self):
        import openpyxl
        with tempfile.TemporaryDirectory() as t:
            folder=Path(t);target=folder/'existing.xlsx'
            wb=openpyxl.Workbook();wb.active.title='Company';wb.active['A1']='keep';wb.save(target);wb.close()
            m.save_sqlite(m.path_bundle(folder)['db'],[price()],[],{'widget_count':1,'found_symbols':[],'warnings':[]})
            m.export_workbook(folder,output_path=target)
            wb=openpyxl.load_workbook(target)
            self.assertEqual(wb['Company']['A1'].value,'keep')
            self.assertEqual(wb['Daily']['D2'].value,14)
            self.assertEqual(wb['RawData']['C2'].value,'14')
            wb.close()

    def test_accumulate_sort_and_preserve_user_sheet(self):
        import openpyxl
        with tempfile.TemporaryDirectory() as t:
            folder=Path(t)
            diag={'widget_count':1,'found_symbols':[],'warnings':[]}
            m.save_sqlite(m.path_bundle(folder)['db'],[price()],[],diag)
            m.export_workbook(folder,today=date(2026,10,17))
            path=m.path_bundle(folder)['xlsx']
            wb=openpyxl.load_workbook(path)
            wb.create_sheet('MyNotes')['A1']='keep me';wb.save(path);wb.close()
            m.save_sqlite(m.path_bundle(folder)['db'],[price(date='2026-10-17',mid='15')],[],diag)
            m.export_workbook(folder,today=date(2026,10,17))
            wb=openpyxl.load_workbook(path)
            self.assertEqual(wb['RawData'].max_row,3)
            self.assertEqual(wb['Daily']['A2'].value,'2026-10-17')
            self.assertEqual(wb['MyNotes']['A1'].value,'keep me')
            self.assertIn('당일',[c.value for c in wb['Latest'][2]])
            wb.close()

    def test_workbook_locked_keeps_old_file_and_db(self):
        with tempfile.TemporaryDirectory() as t:
            folder=Path(t);path=m.path_bundle(folder)['xlsx']
            diag={'widget_count':1,'found_symbols':[],'warnings':[]}
            m.save_sqlite(m.path_bundle(folder)['db'],[price()],[],diag)
            m.export_workbook(folder)
            before=path.read_bytes()
            with patch.object(m.os,'replace',side_effect=PermissionError('Excel lock')):
                with self.assertRaises(PermissionError):m.export_workbook(folder)
            self.assertEqual(before,path.read_bytes())
            self.assertEqual(len(m.load_prices(folder)),1)

    def test_monthly_and_forecast_not_daily(self):
        import openpyxl
        with tempfile.TemporaryDirectory() as t:
            folder=Path(t)
            diag={'widget_count':1,'found_symbols':[],'warnings':[]}
            m.save_sqlite(m.path_bundle(folder)['db'],[price(price_mode='Forecast'),
                price(price_mode='Monthly Average',period_label='1 Oct - 15 Oct 2026')],[],diag)
            m.export_workbook(folder)
            wb=openpyxl.load_workbook(m.path_bundle(folder)['xlsx'])
            self.assertEqual(wb['RawData'].max_row,3)
            self.assertEqual(wb['Daily'].max_row,1)
            self.assertEqual(wb['Latest'].max_row,1)
            wb.close()

if __name__=='__main__':unittest.main()
