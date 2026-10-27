# -*- coding: utf-8 -*-
import tempfile,unittest,sqlite3
from contextlib import closing
from decimal import Decimal
from pathlib import Path
from smbs_rates import (FXQuote,FXPageError,decode_date,decode_difference,parse_row,
                        parse_table_records,save_rates,load_rates,prepare_excel_series,add_summary_sheets)

class SMBSContracts(unittest.TestCase):
    def test_user_supplied_html_visible_cells(self):
        q=parse_row(['2026.09.28','미국 달러 (USD)','1,352.00','8.00','0','0','0','0','0','0'],
                    currency='USD',alt='환율하락 아이콘',classes='ico ico_down')
        self.assertEqual((q.quote_date,q.currency,q.rate,q.difference,q.unit),
                         ('2026-09-28','USD','1352.00','-8.00',1))
    def test_2026_oct_19_partial_zero_columns_dont_replace_base_rate(self):
        q=parse_row(['2026.10.19','미국 달러 (USD)','1,342.40','6.10','0.00','0.00','0.00'],
                    currency='USD',alt='환율하락 아이콘')
        self.assertEqual(q.rate,'1342.40')
        self.assertEqual(q.difference,'-6.10')
    def test_wrong_currency_not_silently_accepted(self):
        with self.assertRaises(FXPageError):
            parse_row(['2026.09.28','미국 달러 (USD)','1,352.00','8.00'],currency='CNH',alt='환율하락 아이콘')
    def test_zero_not_published(self):
        with self.assertRaises(FXPageError):
            parse_row(['2026.09.28','미국 달러 (USD)','0.00','8.00'],currency='USD',alt='환율하락 아이콘')
    def test_diff_without_sign_is_error(self):
        with self.assertRaises(FXPageError):
            decode_difference('8.00','')
    def test_duplicate_render_is_not_duplicate_date(self):
        self.assertEqual(decode_date('2026.09.28 2026.09.28'),'2026-09-28')
    def test_ambiguous_number_is_error(self):
        with self.assertRaises(FXPageError):
            parse_row(['2026.09.28','미국 달러 (USD)','1,352.00 1,360.00','8.00'],currency='USD',alt='환율하락 아이콘')
    def test_repeated_collection_no_duplicates_revisions_recorded(self):
        with tempfile.TemporaryDirectory() as d:
            db=Path(d)/'test.sqlite3'
            q=FXQuote('2026-09-28','USD','1352.00','-8.00',1)
            self.assertEqual(save_rates(db,[q])['inserted'],1)
            self.assertEqual(save_rates(db,[q])['unchanged'],1)
            self.assertEqual(save_rates(db,[FXQuote('2026-09-28','USD','1351.50','-8.50',1)])['updated'],1)
            with closing(sqlite3.connect(db)) as cx, cx:
                self.assertEqual(cx.execute('SELECT COUNT(*) FROM fx_rates').fetchone()[0],1)
                self.assertEqual(cx.execute('SELECT COUNT(*) FROM fx_rate_revisions').fetchone()[0],1)
            self.assertEqual(load_rates(db)[0].rate,'1351.50')
    def test_year_crossing_months_and_stable_daily_keys(self):
        rows=[FXQuote('2026-12-31','JPY','910.00','-2.0',100),
              FXQuote('2027-01-02','JPY','915.00','5.0',100)]
        dates=prepare_excel_series([],rows)
        self.assertEqual(sorted(dates),['2026-12-31','2027-01-02'])
        self.assertEqual(dates['2026-12-31']['FX_JPY_DIFF_KRW_per_100'],-2)
        from openpyxl import Workbook
        wb=Workbook();wb.active.title='원래내시트'
        add_summary_sheets(wb,[],rows)
        self.assertIn('원래내시트',wb.sheetnames)
        self.assertEqual(wb['일별시세'].max_row,4)
        self.assertEqual(wb['월별시세'].max_row,4)
        wb.close()
    def test_user_alias_preserved_on_regeneration(self):
        from openpyxl import Workbook
        wb=Workbook();wb.active.title='내분석'
        q=FXQuote('2026-12-31','USD','1380.00','2.0',1)
        add_summary_sheets(wb,[],[q])
        wb['가격열설정']['B2']='회사환율_USD'
        add_summary_sheets(wb,[],[q,FXQuote('2027-01-02','USD','1390.00','10.0',1)])
        self.assertEqual(wb['일별시세']['B1'].value,'회사환율_USD')
        self.assertEqual(wb['일별시세'].max_row,4)
        self.assertEqual(wb['일별시세'].max_column,2)
        self.assertNotIn('전일대비',str(wb['일별시세']['B1'].value))
        self.assertIsNone(wb['월별시세']['C3'].value)
        wb.close()

    def test_mid_not_invented(self):
        from types import SimpleNamespace
        p=SimpleNamespace(price_mode='Actual',period_label='',date='2026-10-15',symbol='MB-X-1',unit='USD/t',view_key='v1',mid=None,low='12',high='14')
        dates=prepare_excel_series([p],[])
        self.assertNotIn('MB-X-1|USD/t|Actual|v1|MID',dates['2026-10-15'])
        self.assertEqual(len(dates['2026-10-15']),2)

if __name__ == '__main__':unittest.main()
