import tempfile
import unittest
from pathlib import Path
from html.parser import HTMLParser
import test_metal_dashboard as fixtures


class DetailsParser(HTMLParser):
    def __init__(self):super().__init__();self.sources=[];self.order=[]
    def handle_starttag(self,tag,attrs):
        attrs=dict(attrs)
        if tag=='details':self.sources.append(attrs)
        if tag in ('details','table'):self.order.append(tag)


class MetalsReportTests(unittest.TestCase):
    def test_all_failures_retained_only_in_bottom_audit_per_user_policy(self):
        from dataclasses import replace
        from report_exports import export_html
        with tempfile.TemporaryDirectory() as t:
            c=replace(fixtures.MetalDashboardTests().context(t),warnings=('prices 저장 자료 없음','fx 저장 자료 없음','단위 안내','FDR 연결 실패: proxy denied'))
            path=Path(t)/'warning.html';export_html(c,path)
            html=path.read_text();body,audit=html.split('<details id="source-basis"',1)
            self.assertNotIn('FDR 연결 실패',body)
            for warning in c.warnings:self.assertIn(warning,audit)

    def test_metal_pptx_condition_column_and_rows_have_readable_space(self):
        from report_exports import export_pptx
        from pptx import Presentation
        from pptx.util import Inches
        with tempfile.TemporaryDirectory() as t:
            context=fixtures.MetalDashboardTests().context(t)
            path=Path(t)/'metals.pptx';export_pptx(context,path)
            table=next(s.table for s in Presentation(path).slides[0].shapes if s.has_table)
            self.assertGreaterEqual(table.columns[len(table.columns)-1].width,Inches(2.5))
            self.assertGreaterEqual(table.rows[1].height,Inches(1.4))

    def test_source_details_last_default_closed_including_print(self):
        from report_exports import export_html
        with tempfile.TemporaryDirectory() as t:
            c=fixtures.MetalDashboardTests().context(t);p=Path(t)/'report.html';export_html(c,p)
            html=p.read_text();parser=DetailsParser();parser.feed(html)
            self.assertTrue(parser.sources)
            self.assertNotIn('open',parser.sources[0])
            self.assertEqual(parser.sources[-1]['id'],'source-basis')
            self.assertGreater(html.index('자료 상태'),html.index('<details id="source-basis"'))
            self.assertNotIn('details::details-content{content-visibility:visible',html)
            self.assertIn('details:not([open]){display:none}',html)
            self.assertEqual(html.count('<script'),1)
            self.assertIn('<script id="report-controls">',html)


if __name__=='__main__':unittest.main()
