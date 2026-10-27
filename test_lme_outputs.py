"""Synthetic LME storage-to-output read-back; no live official price claims."""
from html import unescape
import json
from pathlib import Path
import tempfile
import unittest


class LMEOutputTests(unittest.TestCase):
    def context(self,folder):
        import FastmarketsCollector_v3_4 as core
        import lme_official
        import report_model
        from test_lme_official import fixture
        # Exercise only the current checkout's modules, never a stale installed copy.
        self.assertEqual(Path(report_model.__file__).resolve().parent,Path(__file__).resolve().parent)
        rows=lme_official.parse(fixture(),'2026-10-27T08:30:00+09:00')['rows']
        self.assertEqual(len(rows),12)
        lme_official.save_rows(core.path_bundle(Path(folder))['db'],rows)
        return report_model.build_context(folder,captured_at='2026-10-27T08:30:00+09:00')

    def test_lme_native_pptx_cells_match_same_html_xlsx_json_snapshot(self):
        import report_exports
        from openpyxl import load_workbook
        from pptx import Presentation
        with tempfile.TemporaryDirectory() as folder:
            context=self.context(folder)
            expected=context.tables['보고용_LME']
            self.assertEqual(len(expected),7)
            self.assertEqual({row[1] for row in expected[1:]},{'2026-10-26'})
            self.assertEqual({row[2] for row in expected[1:]},{'Cash','3-month'})
            self.assertEqual({row[3] for row in expected[1:]},{'Ask'})
            self.assertEqual({row[5] for row in expected[1:]},{'USD/t'})
            result=report_exports.export_reports(context,Path(folder)/'outputs',include_excel=True)
            self.assertEqual(set(result),{'html','pptx','context','excel'})
            tables=[shape.table for slide in Presentation(result['pptx']).slides
                    for shape in slide.shapes if shape.has_table]
            matching=[table for table in tables if [cell.text for cell in table.rows[0].cells]==list(expected[0])]
            self.assertEqual(len(matching),1,'LME prices must appear as one native editable PPTX table')
            actual=[[cell.text for cell in row.cells] for row in matching[0].rows]
            self.assertEqual(actual,[[report_exports.display_value(value) for value in row] for row in expected])
            workbook=load_workbook(result['excel'],read_only=True)
            try:self.assertEqual(list(workbook['보고용_LME'].values),list(expected))
            finally:workbook.close()
            html=unescape(Path(result['html']).read_text())
            self.assertIn('보고용_LME',html)
            for row in expected:
                for value in row:self.assertIn(report_exports.display_value(value),html)
            payload=json.loads(Path(result['context']).read_text())
            self.assertEqual(payload['tables']['보고용_LME'],[list(row) for row in expected])
            self.assertEqual(payload['snapshot_id'],context.snapshot_id)

    def test_existing_report_order_preserved_then_lme_added_without_summary(self):
        from report_model import REPORT_NAMES
        from report_exports import export_pptx
        from pptx import Presentation
        with tempfile.TemporaryDirectory() as folder:
            context=self.context(folder)
            target=Path(folder)/'report.pptx';export_pptx(context,target)
            titles=[shape.text.split('  ')[0] for slide in Presentation(target).slides
                    for shape in slide.shapes if shape.has_text_frame and '  1/' in shape.text]
            self.assertEqual(titles,[name for name in REPORT_NAMES if name!='보고용_종합']+['보고용_LME'])


if __name__=='__main__':unittest.main()
