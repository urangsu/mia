"""Offline report renderers. No browser, AI, news, or network calls."""
from __future__ import annotations
from pathlib import Path
from datetime import datetime
import os
import tempfile
from decimal import Decimal
from report_model import canonical_json,make_workbook_sheets,REPORT_NAMES

def display_value(value):
    if value is None:return '—'
    if isinstance(value,(int,float)):
        rendered=format(Decimal(str(value)),',f')
        return rendered.rstrip('0').rstrip('.') if '.' in rendered else rendered
    return str(value)

def _atomic(path,writer):
    path=Path(path)
    fd,name=tempfile.mkstemp(prefix=path.stem+'.',suffix=path.suffix,dir=path.parent)
    os.close(fd)
    temporary=Path(name)
    try:
        writer(temporary)
        os.replace(temporary,path)
    finally:
        if temporary.exists():temporary.unlink()

def export_html(context,path):
    from jinja2 import Environment,FileSystemLoader,select_autoescape
    env=Environment(loader=FileSystemLoader(Path(__file__).parent),autoescape=select_autoescape(['html']))
    tables=[]
    for name,rows in context.tables.items():
        body=[]
        for row in rows[1:]:
            cells=[]
            for header,value in zip(rows[0],row):
                css='numeric' if isinstance(value,(int,float)) else 'source' if header=='출처' else ''
                if header in ('전일대비','등락률(%)') and isinstance(value,(int,float)):
                    css+=' positive' if value>0 else ' negative' if value<0 else ''
                cells.append(dict(text=display_value(value),css=css))
            body.append(cells)
        tables.append(dict(name=name,headers=rows[0],rows=body))
    failures=[w for w in context.warnings if any(k in w.casefold() for k in ('failed','실패','충돌','철회','mismatch','auth'))]
    notices=[w for w in context.warnings if w not in failures]
    top_warnings=failures+notices[:max(0,3-len(failures))]
    rendered=env.get_template('report_template.html').render(context=context,tables=tables,top_warnings=top_warnings)
    _atomic(path,lambda p:p.write_text(rendered,encoding='utf-8'))
    return str(path)

def export_pptx(context,path):
    from pptx import Presentation
    from pptx.util import Inches,Pt
    from pptx.dml.color import RGBColor
    from pptx.enum.text import PP_ALIGN,MSO_ANCHOR
    prs=Presentation();prs.slide_width=Inches(16);prs.slide_height=Inches(9)
    bg=RGBColor.from_string('101013');fg=RGBColor.from_string('F3EEEE');muted=RGBColor.from_string('AAAAB2')
    def text(slide,value,x,y,w,h,size=14,color=fg):
        shape=slide.shapes.add_textbox(Inches(x),Inches(y),Inches(w),Inches(h))
        shape.text_frame.word_wrap=True
        p=shape.text_frame.paragraphs[0];p.text=value;p.font.size=Pt(size);p.font.color.rgb=color;p.font.name='Malgun Gothic'
        return shape
    # Every table remains native/editable. Eight-row pages leave room for wrapped
    # units/source URLs and the two distinct closing dates, without tiny fonts.
    ordered=[name for name in REPORT_NAMES if name in context.tables]
    ordered.extend(name for name in context.tables if name not in REPORT_NAMES)
    for name in ordered:
        if name=='보고용_종합':continue
        rows=context.tables[name];body=rows[1:]
        metal=name=='보고용_금속시세'
        page_size=3 if metal else 8
        pages=[body[i:i+page_size] for i in range(0,len(body),page_size)] or [()]
        for index,page in enumerate(pages,1):
            slide=prs.slides.add_slide(prs.slide_layouts[6]);slide.background.fill.solid();slide.background.fill.fore_color.rgb=bg
            text(slide,context.label,0.55,.3,14.8,.4,14,muted)
            text(slide,f'{name}  {index}/{len(pages)}',.55,.88,14.8,.55,26)
            table=slide.shapes.add_table(len(page)+1,len(rows[0]),Inches(.55),Inches(1.65),Inches(14.9),Inches(.65*(len(page)+1))).table
            # Source URLs wrap, leaving ample room for dates and numeric columns.
            widths_by_header={'출처':3.6,'단위':1.3,'상태':1.7,'지표':1.4,'종목명':1.8,'시장일':1.5,'비교일':1.5,'고시일':1.5,'순위':.8,'코드':1.3}
            if metal:
                widths_by_header.update({'품목명':2.1,'고시일':1.3,'Mid':1,'구분':1.1,'출처':2.2,'그룹':.9,'원본 조건':4})
                table.rows[0].height=Inches(.55)
                for row in list(table.rows)[1:]:row.height=Inches(1.6)
            widths=[widths_by_header.get(header,1.5) for header in rows[0]]
            scale=14.9/sum(widths)
            for i,w in enumerate(widths):table.columns[i].width=Inches(w*scale)
            for ridx,row in enumerate([rows[0]]+list(page)):
                for cidx,value in enumerate(row):
                    cell=table.cell(ridx,cidx);cell.text=display_value(value)
                    cell.margin_left=Inches(.08);cell.margin_right=Inches(.08);cell.margin_top=Inches(.06);cell.margin_bottom=Inches(.03)
                    cell.vertical_anchor=MSO_ANCHOR.MIDDLE;cell.fill.solid();cell.fill.fore_color.rgb=RGBColor.from_string('202025' if ridx==0 else '17171C')
                    for p in cell.text_frame.paragraphs:
                        p.font.name='Malgun Gothic';p.font.size=Pt(10 if metal and rows[0][cidx] in ('원본 조건','출처') else 13 if ridx else 12);p.font.color.rgb=fg if ridx else muted
                        if isinstance(value,(int,float)):p.alignment=PP_ALIGN.RIGHT
                        if ridx and rows[0][cidx] in ('전일대비','등락률(%)') and isinstance(value,(int,float)) and value:
                            p.font.color.rgb=RGBColor.from_string('F5445A' if value>0 else '4391FF')
            if not page:text(slide,'미고시 · 검증된 관측 자료 없음',.7,2.4,14,.6,18,muted)
            dates=' · '.join(k+': '+((v[0]+' ~ '+v[-1]) if len(v)>1 else v[0] if v else '미고시') for k,v in context.source_dates.items())
            text(slide,dates,.55,7.8,14.9,.35,10,muted)
            text(slide,'자료 기준 '+(context.as_of or '미확인')+' | SHA-256 '+context.snapshot_id,.55,8.2,14.9,.42,10,muted)
    warnings=context.warnings
    for offset in range(0,len(warnings),8):
        slide=prs.slides.add_slide(prs.slide_layouts[6]);slide.background.fill.solid();slide.background.fill.fore_color.rgb=bg
        text(slide,context.label,.55,.3,14.8,.4,14,muted)
        text(slide,'자료 상태',.55,.88,14.8,.55,26)
        for i,warning in enumerate(warnings[offset:offset+8]):text(slide,warning,.65,1.8+i*.68,14.4,.64,16)
        text(slide,'자료 기준 '+(context.as_of or '미확인')+' | SHA-256 '+context.snapshot_id,.55,8.2,14.9,.42,10,muted)
    _atomic(path,prs.save)
    return str(path)

def export_excel(context,path,*,accumulation=False):
    if accumulation:
        from report_model import make_accumulation_sheets
        sheets=make_accumulation_sheets(context)
    else:sheets=make_workbook_sheets(context)
    if os.name=='nt':
        from excel_secure import write_com_sheets
        write_com_sheets(Path(path),sheets)
    else:
        from openpyxl import Workbook
        from openpyxl.styles import Font,PatternFill,Alignment
        from openpyxl.utils import get_column_letter
        wb=Workbook();wb.remove(wb.active)
        for name,rows in sheets.items():
            ws=wb.create_sheet(name)
            for row in rows:ws.append(row)
            ws.freeze_panes='B2';ws.sheet_view.showGridLines=False
            ws.auto_filter.ref=ws.dimensions
            for cell in ws[1]:cell.font=Font(color='FFFFFF',bold=True);cell.fill=PatternFill('solid',fgColor='202025')
            for row in ws.iter_rows(min_row=2):
                for cell in row:
                    if isinstance(cell.value,(int,float)):cell.number_format='#,##0.##########';cell.alignment=Alignment(horizontal='right')
            for i in range(1,ws.max_column+1):ws.column_dimensions[get_column_letter(i)].width=22
        _atomic(path,wb.save)
    return str(path)

def export_reports(context,output_dir,*,include_excel=False):
    """Keep each successful artifact even when another renderer/Excel fails.

    Successful calls retain the existing format:path contract. Partial calls add
    an errors mapping; callers must surface it separately from available paths.
    Only a run with no successful format raises.
    """
    output_dir=Path(output_dir);output_dir.mkdir(parents=True,exist_ok=True)
    prefix='report_'+context.snapshot_id[:16]
    jobs=[('html',output_dir/(prefix+'.html'),lambda path:export_html(context,path)),
          ('pptx',output_dir/(prefix+'.pptx'),lambda path:export_pptx(context,path)),
          ('context',output_dir/(prefix+'_context.json'),
           lambda path:_atomic(path,lambda p:p.write_text(canonical_json(context.to_dict()),encoding='utf-8')))]
    if include_excel:jobs.append(('excel',output_dir/(prefix+'.xlsx'),lambda path:export_excel(context,path)))
    result={};errors={}
    for kind,path,write in jobs:
        try:
            write(path)
            result[kind]=str(path)
        except Exception as exc:
            errors[kind]=f'{type(exc).__name__}: {str(exc)[:300]}'
    if not result:
        raise RuntimeError('보고서 전체 형식 출력 실패: '+'; '.join(kind+': '+error for kind,error in errors.items()))
    if errors:result['errors']=errors
    return result
